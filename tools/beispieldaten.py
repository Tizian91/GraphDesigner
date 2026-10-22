# Erzeugt die Beispieldaten in samples/ (feste Zufallszahlen, damit die Dateien reproduzierbar sind).
import numpy as np, datetime as dt
from openpyxl import Workbook
rng = np.random.default_rng(42)
de = lambda v, d=3: f'{v:.{d}f}'.replace('.', ',')

# messreihe.csv: Semikolon, Dezimalkomma
t = np.arange(50) * 0.2
a = 20 + 5 * (1 - np.exp(-t / 2)) + rng.normal(0, .15, 50)
b = 20 + 3 * np.sin(t) + rng.normal(0, .15, 50)
c = 18 + 0.8 * t + rng.normal(0, .3, 50)
with open('samples/messreihe.csv', 'w', encoding='utf-8') as f:
    f.write('Zeit [s];Sensor A;Sensor B;Sensor C\n')
    for r in zip(t, a, b, c): f.write(';'.join([de(r[0], 1)] + [de(v) for v in r[1:]]) + '\n')

# gruppen.csv: lange Form, Wiederholungen je Zeitpunkt
with open('samples/gruppen.csv', 'w', encoding='utf-8') as f:
    f.write('Zeit,Wert,Gruppe,Bedingung\n')
    for g, (k, base) in {'A': (0.6, 10), 'B': (0.9, 12), 'C': (0.3, 9)}.items():
        for cond, fac in {'Kontrolle': 1.0, 'Behandlung': 1.35}.items():
            for z in range(0, 11, 2):
                for _ in range(6):
                    f.write(f'{z},{base + k * fac * z + rng.normal(0, 1.2):.2f},{g},{cond}\n')

# umsatz.md: Markdown, ein negativer Wert
mon = ['Jan', 'Feb', 'Mär', 'Apr', 'Mai', 'Jun']
rows = [[120, 80, 45, 15], [135, 72, 50, -8], [128, 90, 48, 12], [150, 95, 60, 20], [160, 88, 52, 18], [172, 101, 66, 25]]
with open('samples/umsatz.md', 'w', encoding='utf-8') as f:
    f.write('Umsatz nach Monat und Sparte (Tsd. €)\n\n| Monat | Hardware | Software | Service | Sonstiges |\n|:--|--:|--:|--:|--:|\n')
    for m, r in zip(mon, rows): f.write(f'| {m} | ' + ' | '.join(str(v) for v in r) + ' |\n')

# blasen.csv: ungefähre, gerundete Beispielwerte
laender = [('Deutschland', 84.5, 52700, 81.0, 'Europa'), ('Frankreich', 68.2, 44400, 82.3, 'Europa'),
    ('Italien', 58.9, 38400, 83.0, 'Europa'), ('Polen', 36.7, 22100, 77.5, 'Europa'), ('Schweden', 10.5, 55500, 83.1, 'Europa'),
    ('USA', 334.9, 80400, 77.4, 'Amerika'), ('Brasilien', 216.4, 10000, 73.4, 'Amerika'), ('Mexiko', 128.5, 13900, 75.0, 'Amerika'),
    ('Kanada', 40.1, 53400, 81.7, 'Amerika'), ('Argentinien', 46.7, 13700, 76.1, 'Amerika'),
    ('China', 1410.7, 12600, 78.6, 'Asien'), ('Indien', 1428.6, 2500, 70.8, 'Asien'), ('Japan', 124.5, 33800, 84.7, 'Asien'),
    ('Indonesien', 277.5, 4900, 71.0, 'Asien'), ('Südkorea', 51.7, 33100, 83.7, 'Asien'),
    ('Nigeria', 223.8, 1600, 54.0, 'Afrika'), ('Ägypten', 112.7, 3500, 70.2, 'Afrika'), ('Südafrika', 60.4, 6000, 62.3, 'Afrika'),
    ('Kenia', 55.1, 2000, 62.7, 'Afrika'), ('Äthiopien', 126.5, 1300, 65.6, 'Afrika'),
    ('Australien', 26.6, 64700, 83.1, 'Ozeanien'), ('Neuseeland', 5.2, 48500, 82.5, 'Ozeanien')]
with open('samples/blasen.csv', 'w', encoding='utf-8') as f:
    f.write('Land;Einwohner (Mio.);BIP pro Kopf (US$);Lebenserwartung (Jahre);Kontinent\n')
    for l in laender: f.write(f'{l[0]};{de(l[1], 1)};{l[2]};{de(l[3], 1)};{l[4]}\n')

# verteilung.tsv: 200 Werte in zwei Gruppen
with open('samples/verteilung.tsv', 'w', encoding='utf-8') as f:
    f.write('Gruppe\tMesswert\n')
    for v in rng.normal(50, 8, 110): f.write(f'Werk Nord\t{v:.2f}\n')
    for v in np.concatenate([rng.normal(58, 5, 70), rng.normal(72, 3, 20)]): f.write(f'Werk Süd\t{v:.2f}\n')

# zeitreihe.xlsx: zwei Tabellenblätter mit Datumswerten
wb = Workbook(); ws = wb.active; ws.title = 'Temperatur'
ws.append(['Datum', 'Kassel', 'Hamburg', 'München'])
d0 = dt.date(2025, 1, 1)
for i in range(0, 365, 3):
    doy = 2 * np.pi * (i - 15) / 365
    ws.append([d0 + dt.timedelta(days=i)] + [round(m - a * np.cos(doy) + rng.normal(0, 2), 1) for m, a in [(9.5, 9), (9.8, 8), (9, 10)]])
for row in ws.iter_rows(min_row=2, max_col=1):
    for cell in row: cell.number_format = 'DD.MM.YYYY'
ws2 = wb.create_sheet('Niederschlag')
ws2.append(['Monat', 'Kassel', 'Hamburg', 'München'])
for m in range(1, 13):
    ws2.append([dt.date(2025, m, 1)] + [round(v) for v in rng.gamma(6, 11, 3)])
for row in ws2.iter_rows(min_row=2, max_col=1):
    for cell in row: cell.number_format = 'MM.YYYY'
wb.save('samples/zeitreihe.xlsx')
print('ok')
